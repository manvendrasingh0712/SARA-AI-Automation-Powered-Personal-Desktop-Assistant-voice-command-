"""
sara.skills.notes_qa
Your class notes as a knowledge base: ask questions, get a summary, or be
quizzed — all answered ONLY from what's in Config.NOTES_FOLDER (default
sara_class_notes/ next to config.py), via sara/core/rag.py's LongTermMemory.

Intents
-------
  notes_qa          "what do my notes say about X" / "notes mein X ke baare mein kya likha"
  notes_summary     "summarize my notes on X" / "X ke notes ka summary do"
  notes_quiz        "quiz me on X" / "chapter 4 ka quiz lo"  (3 Q&A, answers on request)
  notes_quiz_answers "quiz answers"                          (within 15 min of a quiz)
  notes_refresh     "refresh my notes" / "notes sync karo"

What gets indexed
-----------------
.txt and .md always; .pdf (pip install pypdf), .docx (pip install python-docx)
and .pptx (pip install python-pptx) automatically when the library is
installed (a one-time log line says what's missing). Chunking is
sentence-aware with overlap (Config.NOTES_CHUNK_CHARS=800,
NOTES_CHUNK_OVERLAP=120) and keeps markdown headings as "[Heading] ..."
context on each chunk.

Sync (sync_notes_folder)
------------------------
Called once at startup by sara/orchestrator/core_wiring.py; also starts a
background WATCHER (Config.NOTES_WATCH_INTERVAL_S, default 60, 0 = off) that
re-syncs when files are added/edited/deleted and retries after a failed sync.
Each file is swapped in ONE operation (LongTermMemory.replace_source_async):
edited notes never leave stale chunks, deleted notes are removed, and a file's
mtime is stamped only AFTER its chunks were really stored (an embedding outage
leaves the old chunks untouched and the file is retried). State is keyed by the
path relative to NOTES_FOLDER, so physics/ch1.md and chem/ch1.md never collide.
A one-time migration (schema "3") re-indexes everything and drops legacy chunks.
Ingestion is asynchronous: sync_notes_folder() returns the number of files QUEUED.

REQUIRES the patched sara/core/rag.py (replace_source_async +
search(source_prefix=...)); with an unpatched rag.py it degrades safely
(only never-indexed files are ingested, a warning is logged).

Answering
---------
Retrieval = vector search restricted to notes at the source (so conversation
memories can't crowd notes out) + a light keyword re-rank (hybrid). The LLM
sees numbered excerpts and must answer only from them. Replies cite the notes
("From your notes: Physics ch3"); the card lists the source chips.
Known limit: brain.generate_response() adds the excerpt prompt to conversation
history (there's no stateless call in the brain API used here).
"""
from __future__ import annotations

import importlib.util
import json
import logging
import os
import re
import threading
import time
from datetime import datetime, timezone

from config import Config

from ._framework import SkillContext, SkillResult, skill

logger = logging.getLogger("sara.skills.notes_qa")

# ── constants ──────────────────────────────────────────────────────────────
_LAST_SYNC_PREF_KEY = "notes_last_sync_at"
_INDEX_PREF_KEY = "notes_index_files"       # JSON list of relative paths currently indexed
_SCHEMA_PREF_KEY = "notes_index_schema"
_SCHEMA_VERSION = "3"                        # bump => one-time full re-index
_MTIME_KEY = "notes_mtime:{rel}"
_QUIZ_WINDOW_S = 900

_sync_lock = threading.Lock()
_index_lock = threading.Lock()
_warned_unpatched = False
_needs_retry = False                         # set when a sync failed; the watcher retries

_TEXT_EXTS = (".txt", ".md")
_OPTIONAL_EXTS = {".pdf": ("pypdf", "PyPDF2"), ".docx": ("docx",), ".pptx": ("pptx",)}
_PIP_NAMES = {".pdf": "pypdf", ".docx": "python-docx", ".pptx": "python-pptx"}
_warned_missing_libs: set = set()


# ═══════════════════════════════════════════════════════════════════════════
# File discovery + reading
# ═══════════════════════════════════════════════════════════════════════════

def _lib_available(names) -> bool:
    return any(importlib.util.find_spec(n) is not None for n in names)


def supported_extensions() -> tuple:
    return _TEXT_EXTS + tuple(e for e, libs in _OPTIONAL_EXTS.items() if _lib_available(libs))


def _iter_note_entries():
    """Yields (absolute_path, relative_posix_path) for every indexable note."""
    folder = getattr(Config, "NOTES_FOLDER", None)
    if not folder or not os.path.isdir(folder):
        return
    exts = supported_extensions()
    for root, _dirs, files in os.walk(folder):
        for fname in sorted(files):
            low = fname.lower()
            if low.endswith(exts):
                path = os.path.join(root, fname)
                yield path, os.path.relpath(path, folder).replace(os.sep, "/")
            else:
                for ext, libs in _OPTIONAL_EXTS.items():
                    if low.endswith(ext) and ext not in _warned_missing_libs:
                        _warned_missing_libs.add(ext)
                        logger.warning(
                            "Found %s notes but none of %s is installed — they are "
                            "skipped. Install with: pip install %s",
                            ext, "/".join(libs), _PIP_NAMES[ext],
                        )


def _iter_note_files():
    """Backward compatible: absolute paths only."""
    for path, _rel in _iter_note_entries():
        yield path


def _read_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader  # type: ignore
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: str) -> str:
    import docx  # python-docx
    document = docx.Document(path)
    lines = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (getattr(para.style, "name", "") or "").lower()
        lines.append(f"## {text}" if style.startswith("heading") else text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _read_pptx(path: str) -> str:
    from pptx import Presentation
    lines = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        lines.append(f"## Slide {n}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        lines.append("")
    return "\n".join(lines)


def read_note_text(path: str):
    """Text of one note, or None if it can't be read (old chunks stay)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return _read_pdf(path)
        if ext == ".docx":
            return _read_docx(path)
        if ext == ".pptx":
            return _read_pptx(path)
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:  # noqa: BLE001
        logger.error("Failed to read %s: %s", path, e)
        return None


# ═══════════════════════════════════════════════════════════════════════════
# Chunking (sentence-aware, with overlap, heading-aware)
# ═══════════════════════════════════════════════════════════════════════════

_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")
_SENT_SPLIT_RE = re.compile(r"(?<=[.!?\u0964])\s+")


def _units(text: str):
    """Yields (heading, sentence) in document order."""
    heading = ""
    para: list = []

    def flush():
        joined = re.sub(r"\s+", " ", " ".join(para)).strip()
        para.clear()
        if not joined:
            return []
        return [(heading, s.strip()) for s in _SENT_SPLIT_RE.split(joined) if s.strip()]

    for line in (text or "").splitlines():
        m = _HEADING_RE.match(line)
        if m:
            yield from flush()
            heading = re.sub(r"\s+", " ", m.group(1)).strip()
        elif not line.strip():
            yield from flush()
        else:
            para.append(line.strip())
    yield from flush()


def _hard_split(sentence: str, limit: int):
    """A sentence longer than `limit` is split on whitespace (never mid-word
    unless a single word exceeds the limit)."""
    if len(sentence) <= limit:
        yield sentence
        return
    current = ""
    for word in sentence.split():
        while len(word) > limit:
            if current:
                yield current
                current = ""
            yield word[:limit]
            word = word[limit:]
        if current and len(current) + 1 + len(word) > limit:
            yield current
            current = word
        else:
            current = f"{current} {word}".strip()
    if current:
        yield current


def chunk_text(text: str, size: int = 800, overlap: int = 120) -> list:
    """
    Splits `text` into chunks of about `size` characters that never cut a
    sentence in half, repeat up to `overlap` characters of the previous chunk
    (so an answer straddling a boundary is still retrievable), and carry the
    current markdown heading as "[Heading] ..." context. Overlap never crosses
    a heading change.
    """
    size = max(80, int(size))
    overlap = max(0, min(int(overlap), size // 2))
    chunks: list = []
    cur: list = []
    cur_len = 0
    fresh = 0
    cur_heading = None

    def prefix(h):
        return f"[{h}] " if h else ""

    def emit():
        if fresh > 0 and cur:
            chunks.append((prefix(cur_heading) + " ".join(cur)).strip())

    for heading, sentence in _units(text):
        if heading != cur_heading:
            emit()
            cur, cur_len, fresh, cur_heading = [], 0, 0, heading
        limit = max(40, size - len(prefix(heading)))
        for piece in _hard_split(sentence, limit):
            if cur and cur_len + 1 + len(piece) > limit:
                if fresh > 0:
                    emit()
                    carry, total = [], 0
                    for s in reversed(cur):
                        if total + len(s) + 1 > overlap:
                            break
                        carry.insert(0, s)
                        total += len(s) + 1
                    cur, cur_len, fresh = carry, total, 0
                    if cur and cur_len + 1 + len(piece) > limit:
                        cur, cur_len = [], 0
                else:  # only carried-over text so far and it doesn't fit: drop it
                    cur, cur_len = [], 0
            cur.append(piece)
            cur_len += len(piece) + 1
            fresh += 1
    emit()
    return chunks


# ═══════════════════════════════════════════════════════════════════════════
# Preferences helpers (failure-tolerant)
# ═══════════════════════════════════════════════════════════════════════════

def _pref_get(db, key):
    if db is None:
        return None
    try:
        return db.get_preference(key)
    except Exception as e:  # noqa: BLE001
        logger.warning("get_preference(%r) failed: %s", key, e)
        return None


def _pref_set(db, key, value) -> None:
    if db is None:
        return
    try:
        db.set_preference(key, value)
    except Exception as e:  # noqa: BLE001
        logger.warning("set_preference(%r) failed: %s", key, e)


def _pref_clear(db, key) -> None:
    if db is None:
        return
    try:
        if hasattr(db, "delete_preference"):
            db.delete_preference(key)
        else:
            db.set_preference(key, "")
    except Exception as e:  # noqa: BLE001
        logger.warning("clear preference %r failed: %s", key, e)


def _load_index(db) -> set:
    raw = _pref_get(db, _INDEX_PREF_KEY)
    if not raw:
        return set()
    try:
        data = json.loads(raw)
        return {str(x) for x in data} if isinstance(data, list) else set()
    except (ValueError, TypeError):
        return set()


def _index_update(db, add=None, remove=None) -> None:
    with _index_lock:
        idx = _load_index(db)
        if add:
            idx.add(add)
        if remove:
            idx.discard(remove)
        _pref_set(db, _INDEX_PREF_KEY, json.dumps(sorted(idx)))


# ═══════════════════════════════════════════════════════════════════════════
# Sync
# ═══════════════════════════════════════════════════════════════════════════

def _read_chunks(path: str, rel: str):
    """Chunks of one file; None if unreadable (old chunks stay searchable)."""
    text = read_note_text(path)
    if text is None:
        return None
    chunks = chunk_text(
        text,
        int(getattr(Config, "NOTES_CHUNK_CHARS", 800)),
        int(getattr(Config, "NOTES_CHUNK_OVERLAP", 120)),
    )
    max_chunks = int(getattr(Config, "NOTES_MAX_CHUNKS_PER_FILE", 200))
    if len(chunks) > max_chunks:
        logger.warning("%s: stopping at %d chunks (NOTES_MAX_CHUNKS_PER_FILE) — file is larger.",
                       rel, max_chunks)
        chunks = chunks[:max_chunks]
    return chunks


class _Batch:
    """Counts finished replace jobs; runs `on_all_ok` when the last one
    finishes and none failed."""

    def __init__(self, on_all_ok):
        self._lock = threading.Lock()
        self._pending = 0
        self._failed = False
        self._sealed = False
        self._on_all_ok = on_all_ok

    def add(self) -> None:
        with self._lock:
            self._pending += 1

    def finished(self, ok: bool) -> None:
        global _needs_retry
        with self._lock:
            self._pending -= 1
            self._failed = self._failed or not ok
            fire = self._sealed and self._pending == 0 and not self._failed
        if not ok:
            _needs_retry = True
        if fire:
            self._on_all_ok()

    def seal(self) -> None:
        with self._lock:
            self._sealed = True
            fire = self._pending == 0 and not self._failed
        if fire:
            self._on_all_ok()


def _submit_replace(rag_memory, source, chunks, batch, on_success) -> bool:
    fut = rag_memory.replace_source_async(source, chunks)
    if fut is None:
        return False
    batch.add()

    def _done(f):
        try:
            exc = f.exception()
        except Exception as e:  # noqa: BLE001
            exc = e
        if exc is None:
            try:
                on_success()
            except Exception:  # noqa: BLE001
                logger.exception("post-index bookkeeping failed for %s", source)
            batch.finished(True)
        else:
            logger.error("%s: re-index failed, old data kept, will retry: %s", source, exc)
            batch.finished(False)

    fut.add_done_callback(_done)
    return True


def _sync_legacy(rag_memory, db) -> int:
    """Unpatched rag.py: ingest only never-indexed files (edits would duplicate)."""
    ingested = 0
    for path, rel in _iter_note_entries():
        key = _MTIME_KEY.format(rel=rel)
        if _pref_get(db, key):
            continue
        chunks = _read_chunks(path, rel)
        if chunks is None:
            continue
        for c in chunks:
            rag_memory.add_memory(c, source=f"notes:{rel}")
        try:
            _pref_set(db, key, str(int(os.path.getmtime(path))))
        except OSError:
            continue
        ingested += 1
    return ingested


def sync_notes_folder(rag_memory, db) -> int:
    """
    Queues every new/changed note for (re-)indexing, queues removal of notes
    that no longer exist, starts the background watcher (once) and returns the
    number of files queued. Safe to call repeatedly. Never raises.
    """
    global _warned_unpatched, _needs_retry

    if rag_memory is None or not getattr(rag_memory, "enabled", False):
        return 0

    start_notes_watcher(rag_memory, db)

    if hasattr(rag_memory, "check_backend") and not rag_memory.check_backend():
        _needs_retry = True
        model = getattr(Config, "EMBEDDING_MODEL", "gemini-embedding-001")
        print(
            f"[NotesQA] Gemini embedding model '{model}' isn't responding — "
            f"notes won't be searchable until it is. Check that GEMINI_API_KEY "
            f"is set in .env, that you're online, and that EMBEDDING_MODEL is "
            f"'gemini-embedding-001' (or unset). Will retry automatically."
        )
        return 0

    with _sync_lock:
        try:
            _needs_retry = False
            if not callable(getattr(rag_memory, "replace_source_async", None)):
                if not _warned_unpatched:
                    _warned_unpatched = True
                    logger.warning(
                        "rag.py has no replace_source_async(): apply patches/rag_patch.py. "
                        "Edited/deleted notes will NOT be re-indexed until then."
                    )
                return _sync_legacy(rag_memory, db)

            force = db is not None and _pref_get(db, _SCHEMA_PREF_KEY) != _SCHEMA_VERSION
            previous = _load_index(db)
            entries = list(_iter_note_entries())
            current = {rel for _p, rel in entries}

            def _all_ok():
                _pref_set(db, _LAST_SYNC_PREF_KEY, datetime.now(timezone.utc).isoformat())

            batch = _Batch(_all_ok)

            if force:  # drop legacy chunks stored under bare filenames FIRST (queue is FIFO)
                for _path, rel in entries:
                    base = os.path.basename(rel)
                    if base != rel:
                        rag_memory.replace_source_async(f"notes:{base}", [])

            queued = 0
            for path, rel in entries:
                try:
                    mtime = str(int(os.path.getmtime(path)))
                except OSError:
                    continue
                key = _MTIME_KEY.format(rel=rel)
                if not force and _pref_get(db, key) == mtime:
                    continue
                chunks = _read_chunks(path, rel)
                if chunks is None:
                    _needs_retry = True
                    continue
                if _submit_replace(
                    rag_memory, f"notes:{rel}", chunks, batch,
                    on_success=lambda k=key, m=mtime, r=rel: (
                        _pref_set(db, k, m), _index_update(db, add=r)
                    ),
                ):
                    queued += 1

            for gone in sorted(previous - current):  # notes deleted from disk
                _submit_replace(
                    rag_memory, f"notes:{gone}", [], batch,
                    on_success=lambda g=gone: (
                        _pref_clear(db, _MTIME_KEY.format(rel=g)),
                        _index_update(db, remove=g),
                    ),
                )

            if force:
                # Safe now: a file whose re-index fails keeps its old chunks (atomic)
                # or has no mtime key, so it is retried.
                _pref_set(db, _SCHEMA_PREF_KEY, _SCHEMA_VERSION)
            if queued:
                print(f"[NotesQA] Queued {queued} note file(s) from "
                      f"{getattr(Config, 'NOTES_FOLDER', '?')} for indexing")
            batch.seal()
            return queued
        except Exception as e:  # noqa: BLE001
            _needs_retry = True
            print(f"[NotesQA] sync_notes_folder failed: {e}")
            return 0


def get_notes_index_status(db) -> dict:
    """{"count": int, "last_synced": ISO-string | None}. Walks the live folder,
    so deleted/renamed files don't inflate the count. Never raises."""
    if db is None:
        return {"count": 0, "last_synced": None}
    count = 0
    try:
        for _path, rel in _iter_note_entries():
            if _pref_get(db, _MTIME_KEY.format(rel=rel)):
                count += 1
    except Exception as e:  # noqa: BLE001
        logger.warning("get_notes_index_status: folder walk failed: %s", e)
    return {"count": count, "last_synced": _pref_get(db, _LAST_SYNC_PREF_KEY)}


# ── watcher ────────────────────────────────────────────────────────────────

_watcher_lock = threading.Lock()
_watcher: dict = {"thread": None, "stop": threading.Event()}


def _folder_signature():
    sig = []
    for path, rel in _iter_note_entries():
        try:
            st = os.stat(path)
            sig.append((rel, st.st_mtime_ns, st.st_size))
        except OSError:
            continue
    return hash(tuple(sorted(sig)))


def start_notes_watcher(rag_memory, db) -> bool:
    """Polls the notes folder every Config.NOTES_WATCH_INTERVAL_S seconds
    (default 60; 0 disables) and re-syncs on change or after a failed sync.
    Idempotent; a daemon thread that exits when the RAG store closes."""
    interval = float(getattr(Config, "NOTES_WATCH_INTERVAL_S", 60))
    if interval <= 0:
        return False
    with _watcher_lock:
        t = _watcher["thread"]
        if t is not None and t.is_alive():
            return False
        stop = threading.Event()
        _watcher["stop"] = stop
        # Captured HERE, synchronously in the caller's thread, before the
        # background thread even exists -- so a file the caller writes right
        # after this call returns is never accidentally folded into the
        # watcher's own baseline (which would silently delay detecting it
        # until the NEXT real change, since threads take a moment to start).
        baseline = _folder_signature()

        def loop():
            last = baseline
            while not stop.wait(interval):
                if getattr(rag_memory, "_closed", False) or not getattr(rag_memory, "enabled", False):
                    break
                try:
                    sig = _folder_signature()
                    if sig != last or _needs_retry:
                        last = sig
                        sync_notes_folder(rag_memory, db)
                except Exception:  # noqa: BLE001
                    logger.exception("notes watcher iteration failed")

        thread = threading.Thread(target=loop, name="notes-watcher", daemon=True)
        _watcher["thread"] = thread
        thread.start()
        return True


def stop_notes_watcher() -> None:
    _watcher["stop"].set()


# ═══════════════════════════════════════════════════════════════════════════
# Retrieval (hybrid) + presentation helpers
# ═══════════════════════════════════════════════════════════════════════════

_STOPWORDS = frozenset(
    "the a an of on in to and or is are was were what who how why when about my me "
    "notes note say says said tell give please kya hai ke ka ki mein main se ko aur "
    "batao likha likhe baare bare".split()
)


def _tokens(text: str) -> list:
    return re.findall(r"[a-z0-9\u0900-\u097f]+", (text or "").lower())


def rerank(query: str, hits: list, top_k: int, alpha: float = 0.25) -> list:
    """Light hybrid: (1-alpha)*vector score + alpha*keyword coverage."""
    q = {t for t in _tokens(query) if len(t) > 2 and t not in _STOPWORDS}

    def kw(h):
        if not q:
            return 0.0
        return len(q & set(_tokens(h.text))) / len(q)

    return sorted(hits, key=lambda h: (1 - alpha) * h.score + alpha * kw(h), reverse=True)[:top_k]


def pretty_source(source: str) -> str:
    rel = source[len("notes:"):] if source.startswith("notes:") else source
    base = re.sub(r"\.[A-Za-z0-9]+$", "", rel)
    label = " ".join(p.replace("_", " ").replace("-", " ") for p in base.split("/")[-2:])
    return re.sub(r"\s+", " ", label).strip() or rel


def _unique_sources(hits) -> list:
    seen, out = set(), []
    for h in hits:
        label = pretty_source(h.source)
        if label not in seen:
            seen.add(label)
            out.append(label)
    return out


def _retrieve(ctx: SkillContext, query: str, top_k: int) -> list:
    rag = ctx.notes_memory
    fetch = top_k * 3
    try:
        try:
            hits = rag.search(query, top_k=fetch, source_prefix="notes:")
        except TypeError:  # unpatched rag.py: no source_prefix
            hits = rag.search(query, top_k=fetch * 2)
    except Exception as e:  # noqa: BLE001
        logger.error("search failed: %s", e)
        return []
    hits = [h for h in hits if str(getattr(h, "source", "")).startswith("notes:")]
    return rerank(query, hits, top_k)


def _llm(ctx: SkillContext, prompt: str):
    brain = ctx.brain
    if brain is None or not hasattr(brain, "generate_response"):
        return None
    try:
        text = brain.generate_response(prompt)
        return text.strip() if text and text.strip() else None
    except Exception as e:  # noqa: BLE001
        logger.error("LLM call failed: %s", e)
        return None


def _lang_rule(ctx: SkillContext) -> str:
    return ("Reply in simple Hinglish (Hindi written in Roman script)."
            if ctx.lang == "hinglish" else "Reply in English.")


def _excerpts(hits) -> str:
    return "\n\n".join(f"[{i}] ({pretty_source(h.source)}) {h.text}" for i, h in enumerate(hits, 1))


def _query(match, ctx: SkillContext) -> str:
    groups = [g for g in (match.groups() if match else ()) if g]
    q = groups[0] if groups else ctx.user_input
    return re.sub(r"\s+", " ", q).strip(" .?!,")


def _ready(ctx: SkillContext):
    """Friendly SkillResult if notes search can't run, else None."""
    rag = ctx.notes_memory
    if rag is None or not getattr(rag, "enabled", False):
        return SkillResult(text=ctx.t(
            "Notes search isn't set up right now — long-term memory is disabled.",
            "Notes search abhi set up nahi hai — long-term memory band hai."))
    return None


# ═══════════════════════════════════════════════════════════════════════════
# Skills
# ═══════════════════════════════════════════════════════════════════════════

@skill(
    name="notes_qa",
    patterns=[
        r"(?:what do |do )?my notes (?:say|mention) (?:about |on )?(.+)",
        r"(?:check|search|look in) my notes (?:for|about) (.+)",
        r"notes (?:mein|main) (.+) (?:ke baare mein|ke bare me) kya (?:likha|hai)",
        r"ask (?:my )?notes (?:about )?(.+)",
    ],
    gate=("notes",),
    description="Answers questions from your class notes (txt/md/pdf/docx/pptx) with sources",
    category="study",
    requires=("notes_memory",),
    examples=("what do my notes say about newton's laws", "check my notes for photosynthesis"),
)
def handle(match, ctx: SkillContext):
    blocked = _ready(ctx)
    if blocked:
        return blocked
    query = _query(match, ctx)
    ctx.status("thinking")
    top_k = int(getattr(Config, "NOTES_QA_TOP_K", 4))
    hits = _retrieve(ctx, query, top_k)
    if not hits:
        return SkillResult(text=ctx.t("I couldn't find anything about that in your notes.",
                                      "Aapke notes mein iske baare mein kuch nahi mila."))

    prompt = (
        "Using ONLY the notes excerpts below, answer the question in 2-3 sentences. "
        "If the excerpts don't actually answer it, say so. "
        f"{_lang_rule(ctx)}\n\nNotes excerpts:\n{_excerpts(hits)}\n\nQuestion: {query}"
    )
    answer = _llm(ctx, prompt) or hits[0].text[:400]
    sources = _unique_sources(hits)
    cite = ctx.t("From your notes: ", "Aapke notes se: ") + ", ".join(sources) + "."
    short = query[:40]
    return SkillResult(
        text=f"{answer} {cite}",
        display=f"{answer}\n\n{ctx.t('Sources', 'Sources')}: {', '.join(sources)}",
        card={"type": "notes_answer", "question": query, "answer": answer,
              "sources": [{"label": pretty_source(h.source), "snippet": h.text[:160]} for h in hits]},
        chips=[f"Quiz me on {short}", f"Summarize my notes on {short}"],
    )


@skill(
    name="notes_summary",
    patterns=[
        r"summari[sz]e my notes (?:on|about) (.+)",
        r"(?:give me )?(?:a )?summary of my notes (?:on|about) (.+)",
        r"(?:my )?notes (?:ka |ki |ke )?summary (?:of |on |about )?(.+)",
        r"(.+?) (?:ke |ka )?notes ka summary (?:do|batao)",
    ],
    gate=("summar",),
    description="Summarizes your notes on a topic in a few bullets",
    category="study",
    requires=("notes_memory",),
    examples=("summarize my notes on thermodynamics",),
)
def handle_summary(match, ctx: SkillContext):
    blocked = _ready(ctx)
    if blocked:
        return blocked
    topic = _query(match, ctx)
    ctx.ack("Let me read through your notes.", "Ek second, notes padh rahi hoon.")
    hits = _retrieve(ctx, topic, 6)
    if not hits:
        return SkillResult(text=ctx.t("I couldn't find notes on that topic.",
                                      "Is topic par mujhe notes nahi mile."))
    prompt = (
        "Summarize the notes excerpts below about the topic in at most 4 short "
        f"bullet points, using ONLY the excerpts. {_lang_rule(ctx)}\n\n"
        f"Topic: {topic}\n\nExcerpts:\n{_excerpts(hits)}"
    )
    summary = _llm(ctx, prompt) or " ".join(h.text[:200] for h in hits[:2])
    sources = _unique_sources(hits)
    spoken = re.sub(r"^\s*[-*\u2022]\s*", "", summary, flags=re.M)
    return SkillResult(
        text=f"{spoken} " + ctx.t("From your notes: ", "Aapke notes se: ") + ", ".join(sources) + ".",
        display=f"{summary}\n\n{ctx.t('Sources', 'Sources')}: {', '.join(sources)}",
        card={"type": "notes_summary", "topic": topic, "summary": summary, "sources": sources},
        chips=[f"Quiz me on {topic[:40]}"],
    )


_quiz_state: dict = {"at": 0.0, "topic": "", "qa": []}
_QA_RE = re.compile(r"Q\d*[:.)]\s*(.+?)\s*\n\s*A\d*[:.)]\s*(.+?)(?=\n\s*Q\d*[:.)]|\Z)", re.S | re.I)


def parse_quiz(raw: str) -> list:
    return [(q.strip(), a.strip()) for q, a in _QA_RE.findall(raw or "") if q.strip() and a.strip()]


@skill(
    name="notes_quiz",
    patterns=[
        r"quiz me (?:on|about) (.+?)(?: from my notes)?\s*$",
        r"(?:my )?notes se (.+) ka quiz (?:lo|do|banao)",
        r"(.+) ka quiz (?:lo|do|banao)",
    ],
    gate=("quiz",),
    description="Asks you 3 quiz questions made from your notes",
    category="study",
    requires=("notes_memory",),
    examples=("quiz me on photosynthesis", "chapter 4 ka quiz lo"),
)
def handle_quiz(match, ctx: SkillContext):
    blocked = _ready(ctx)
    if blocked:
        return blocked
    topic = _query(match, ctx)
    ctx.ack("Making a quiz from your notes.", "Notes se quiz bana rahi hoon.")
    hits = _retrieve(ctx, topic, 5)
    if not hits:
        return SkillResult(text=ctx.t("I couldn't find notes on that topic to quiz you.",
                                      "Is topic par quiz ke liye notes nahi mile."))
    raw = _llm(ctx, (
        "From ONLY the notes excerpts below, write exactly 3 short quiz questions with "
        "brief answers. Use exactly this format and nothing else:\n"
        "Q1: ...\nA1: ...\nQ2: ...\nA2: ...\nQ3: ...\nA3: ...\n"
        f"{_lang_rule(ctx)}\n\nExcerpts:\n{_excerpts(hits)}"
    ))
    qa = parse_quiz(raw or "")[:3]
    if not qa:
        return SkillResult(text=ctx.t("I couldn't make a quiz from that right now.",
                                      "Abhi is par quiz nahi bana paayi."))
    _quiz_state.update(at=time.time(), topic=topic, qa=qa)
    numbered = " ".join(f"{i}: {q}" for i, (q, _a) in enumerate(qa, 1))
    return SkillResult(
        text=ctx.t(f"Here are {len(qa)} questions. {numbered} Say 'quiz answers' when you're ready.",
                   f"Ye rahe {len(qa)} sawaal. {numbered} Jab taiyaar ho to 'quiz answers' bolo."),
        card={"type": "quiz", "topic": topic,
              "cards": [{"question": q, "answer": a} for q, a in qa]},
        chips=[ctx.t("Quiz answers", "Quiz answers")],
    )


@skill(
    name="notes_quiz_answers",
    patterns=[
        r"^\s*(?:the )?(?:quiz )?answers?(?: batao| do| bolo| please)?\s*[.!]*\s*$",
        r"^\s*quiz ke answers?(?: batao| do)?\s*[.!]*\s*$",
    ],
    gate=("answer",),
    description="Reads out the answers to the quiz you were just given",
    category="study",
)
def handle_quiz_answers(match, ctx: SkillContext):
    if not _quiz_state["qa"] or time.time() - _quiz_state["at"] > _QUIZ_WINDOW_S:
        return None  # no recent quiz: not for us
    parts = " ".join(f"{i}: {a}" for i, (_q, a) in enumerate(_quiz_state["qa"], 1))
    return SkillResult(
        text=ctx.t(f"The answers: {parts}", f"Answers ye rahe: {parts}"),
        chips=[f"Quiz me on {_quiz_state['topic'][:40]}"],
    )


@skill(
    name="notes_refresh",
    patterns=[
        r"(?:refresh|sync|re-?index|update) (?:my )?notes",
        r"notes (?:refresh|sync|update) karo",
        r"notes (?:ko )?(?:refresh|sync|update) karo",
    ],
    gate=("notes",),
    description="Re-scans your notes folder and indexes new or changed files",
    category="study",
    requires=("notes_memory",),
    examples=("refresh my notes", "notes sync karo"),
)
def handle_refresh(match, ctx: SkillContext):
    blocked = _ready(ctx)
    if blocked:
        return blocked
    ctx.ack("Refreshing your notes.", "Notes refresh kar rahi hoon.")
    queued = sync_notes_folder(ctx.notes_memory, ctx.db)
    status = get_notes_index_status(ctx.db)
    if queued:
        text = ctx.t(f"Queued {queued} note file{'s' if queued != 1 else ''} for indexing. "
                     f"{status['count']} already searchable.",
                     f"{queued} note file indexing ke liye lagayi hai. "
                     f"{status['count']} pehle se searchable hain.")
    else:
        text = ctx.t(f"Your notes are up to date — {status['count']} files indexed.",
                     f"Aapke notes up to date hain — {status['count']} files indexed hain.")
    return SkillResult(text=text, card={"type": "notes_status", **status, "queued": queued})